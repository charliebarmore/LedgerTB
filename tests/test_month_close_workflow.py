"""250 source rows through actual review UI, encrypted accounting and recovery."""
from io import BytesIO
import json
import re

import openpyxl
import pypdfium2 as pdfium
import pytest
import streamlit as st
from streamlit.testing.v1 import AppTest

from database import connection as dbconn
from models import close_map
from models.audit_log import AuditLog
from models.fiscal_period import FiscalPeriod
from models.journal_entry import JournalEntry
from models.reconciliation import BankReconciliation
from models.reports import ReportGenerator
from models.transaction import ImportedTransaction
from money import to_cents
from services import import_review_drafts as drafts, mcp_tools
from services.backups import create_backup, restore_backup
from services.close_package import build_close_package, build_close_package_pdf, load_close_package_snapshot
from services.import_identity import classify_import_duplicates
from tests.conftest import page_path, post_entry
from tests.helpers.month_close import BALANCES, JANUARY, create_month, review_rows
from utils.import_review import row_key


def _button(at, label):
    return next(b for b in at.button if b.label == label)


def _assert_close(client_id):
    tb = mcp_tools.trial_balance(client_id, '2026-01-31')
    assert {r['number']: (to_cents(r['debit']), to_cents(r['credit'])) for r in tb['accounts']} == BALANCES
    assert to_cents(tb['total_debits']) == to_cents(tb['total_credits']) == 1_760_000
    income = mcp_tools.income_statement(client_id, '2026-01-01', '2026-01-31')
    assert tuple(to_cents(income[k]) for k in ('total_revenue', 'total_expenses', 'net_income')) == (1_000_000, 247_500, 752_500)
    bs = mcp_tools.balance_sheet(client_id, '2026-01-31')
    assert tuple(to_cents(bs[k]) for k in ('total_assets', 'total_liabilities', 'total_equity')) == (1_506_500, 160_000, 1_346_500)
    assert bs['balanced']


def test_month_close_partial_post_restart_reconcile_export_restore(db, monkeypatch, tmp_path, fake_credential_vault):
    import utils.client_selector as selector
    client_id, other_id, accounts = create_month()
    monkeypatch.setattr(selector, 'render_client_selector', lambda: client_id)
    monkeypatch.setattr(st, 'page_link', lambda *a, **kw: None)
    fake_credential_vault['categorization_provider'] = 'off'
    post_entry(client_id, JANUARY[0], [(accounts['cash'], 5000, 0), (accounts['capital'], 0, 5000)], entry_type='Beginning Balance')
    rows = review_rows(accounts)
    assert len(rows) == 250
    assert classify_import_duplicates(rows, client_id) == 2
    for row in rows:
        row['batch_id'] = 'juniper-january'
    assert sum(r.get('include', True) for r in rows) == 244
    assert sum(not r.get('selected_account_id') for r in rows) == 3
    revision = drafts.save(client_id, rows)

    # A new window resumes the durable checkpoint, not another window's memory.
    at = AppTest.from_file(page_path('pages/4_Import_Transactions.py'), default_timeout=120)
    at.session_state['import_active_tab'] = 'Review & Categorize'
    at.run()
    _button(at, 'Resume saved review').click().run()
    assert not at.exception and len(at.session_state['transactions_to_review']) == 250
    assert sum(r.get('include', True) for r in at.session_state['transactions_to_review']) == 244

    # Fail AFTER the entry has been written on its transaction connection.
    original_save = ImportedTransaction.save
    failed = rows[0]['description']
    def fail_once(self, *args, **kwargs):
        if self.description == failed:
            raise OSError('Synthetic interruption before import record commit')
        return original_save(self, *args, **kwargs)
    monkeypatch.setattr(ImportedTransaction, 'save', fail_once)
    _button(at, 'Post Transactions').click().run()
    assert not at.exception
    assert JournalEntry.count(client_id) == 241  # opening + 240 complete imported rows
    assert len(ImportedTransaction.get_by_status(client_id, 'Posted')) == 240
    assert len(at.session_state['transactions_to_review']) == 4
    with dbconn.get_cursor() as cur:
        assert cur.execute('SELECT COUNT(*) FROM journal_entries WHERE description=?', (failed,)).fetchone()[0] == 0
    monkeypatch.setattr(ImportedTransaction, 'save', original_save)
    _button(at, 'Post Transactions').click().run()
    assert JournalEntry.count(client_id) == 242
    assert len(at.session_state['transactions_to_review']) == 3
    _button(at, 'Save review for later').click().run()
    saved_revision, unresolved = drafts.load(client_id)
    assert saved_revision != revision and len(unresolved) == 3

    # Competing window must not overwrite this new revision.
    with pytest.raises(drafts.ReviewConflict):
        drafts.save(client_id, rows, expected_revision=revision)
    key = dbconn.get_active_key()
    dbconn.clear_active_key()
    dbconn.set_active_key(key)
    dbconn.init_database()
    resumed = AppTest.from_file(page_path('pages/4_Import_Transactions.py'), default_timeout=120)
    resumed.session_state['import_active_tab'] = 'Review & Categorize'
    resumed.run()
    _button(resumed, 'Resume saved review').click().run()
    assert not resumed.exception and len(resumed.session_state['transactions_to_review']) == 3
    assert all(not r.get('selected_account_id') for r in resumed.session_state['transactions_to_review'])
    # The conservative journal matcher also flags a same-day $100 purchase
    # beside a $100 customer receipt. Human evidence confirms these differ.
    flagged = [r for r in resumed.session_state['transactions_to_review'] if r.get('is_duplicate')]
    assert len(flagged) == 1 and flagged[0]['duplicate_kind'] == 'journal_match'
    resumed.checkbox(key=row_key('duplicate_override', flagged[0])).check().run()
    assert all(r.get('include', True) for r in resumed.session_state['transactions_to_review'])
    # Human resolves missing evidence and chooses gross accounts; separate AJEs
    # below allocate the documented personal portions of the mixed purchases.
    for row in resumed.session_state['transactions_to_review']:
        if 'missing receipt' in row['description']:
            row['receipt_text'] = 'Fictional receipt: business equipment repair $75, no personal portion.'
    resumed.run()  # Evidence changes invalidate an earlier category selection.
    for row in resumed.session_state['transactions_to_review']:
        if 'missing receipt' in row['description']:
            target = accounts['repairs']
        else:
            target = accounts['office'] if 'office and personal' in row['description'] else accounts['software']
        resumed.selectbox(key=row_key('cat', row)).select(target)
    resumed.run()
    _button(resumed, 'Post Transactions').click().run()
    assert not resumed.exception and JournalEntry.count(client_id) == 245, (resumed.session_state['transactions_to_review'], [w.value for w in resumed.warning], [e.value for e in resumed.error])
    assert len(ImportedTransaction.get_by_status(client_id, 'Posted')) == 244
    assert not resumed.session_state['transactions_to_review']

    # Re-import detects all posted source rows plus the two duplicate receipts.
    repeated = review_rows(accounts)
    assert classify_import_duplicates(repeated, client_id) == 246
    for expense, personal in [('office', 40), ('software', 20)]:
        post_entry(client_id, JANUARY[1], [(accounts['draw'], personal, 0), (accounts[expense], 0, personal)],
                   entry_type='Adjusting', source_reference='Juniper mixed purchase allocation')
    post_entry(client_id, JANUARY[1], [(accounts['wages'], 300, 0), (accounts['accrual'], 0, 300)],
               entry_type='Adjusting', source_reference='Juniper unpaid January wages')
    assert JournalEntry.count(client_id) == 248
    _assert_close(client_id)
    assert JournalEntry.count(other_id) == 0

    reconciliations = []
    for source, ending, count in [('cash', 15065, 149), ('card', 1300, 100)]:
        reconciliation = BankReconciliation.create(client_id, accounts[source], *JANUARY, ending)
        lines = reconciliation.lines()
        assert len(lines) == count
        reconciliation.save_selected_lines([line.line_id for line in lines])
        assert reconciliation.difference() == 0
        reconciliation.complete()
        reconciliations.append(reconciliation.id)
    year = next(p for p in FiscalPeriod.get_all(client_id) if p.period_type == 'Year')
    for account_id in accounts.values():
        close_map.save_explanation(client_id, year.id, account_id, 'Agrees to independently specified Juniper January schedule.')
        close_map.add_evidence(client_id, year.id, account_id, 'workpaper', 'Juniper-Jan', 'Fictional monthly close')
        close_map.signoff(client_id, year.id, account_id, 'preparer')
        close_map.signoff(client_id, year.id, account_id, 'reviewer')
    assert close_map.readiness(client_id, year.id)['ready']

    tb_rows, _ = ReportGenerator.trial_balance_worksheet(client_id, *JANUARY)
    snapshot = load_close_package_snapshot(client_id, *JANUARY)
    workbook = build_close_package(client_id, 'Juniper Month Services', *JANUARY, tb_rows, snapshot=snapshot)
    wb = openpyxl.load_workbook(BytesIO(workbook.getvalue()))
    income_cells = {r[0]: r[1] for r in wb['Income Statement'].iter_rows(values_only=True)}
    assert income_cells['NET INCOME'] == 7525
    assert income_cells['Total Expenses'] == 2475
    summary = {r[0]: r[1] for r in wb['Summary'].iter_rows(values_only=True)}
    assert summary['Final trial balance - total debits'] == 17600
    assert wb['Transactions'].max_row == 497
    pdf = build_close_package_pdf(client_id, 'Juniper Month Services', *JANUARY, tb_rows, snapshot=snapshot)
    (tmp_path / 'juniper-close.xlsx').write_bytes(workbook.getvalue())
    (tmp_path / 'juniper-close.pdf').write_bytes(pdf.getvalue())
    with pdfium.PdfDocument(pdf.getvalue()) as document:
        text = '\n'.join(page.get_textpage().get_text_range() for page in document)
    assert re.search(r'NET INCOME\s+7,525\.00', text)
    assert re.search(r'Total Assets\s+15,065\.00', text)

    backup = create_backup(tmp_path / 'backups')
    assert backup.database_path.read_bytes()[:16] != b'SQLite format 3\x00'
    prior_count = JournalEntry.count(client_id)
    post_entry(client_id, JANUARY[1], [(accounts['cash'], 1, 0), (accounts['revenue'], 0, 1)])
    assert not close_map.readiness(client_id, year.id)['ready']
    def restore_audit(conn):
        AuditLog.write(conn.cursor(), None, 'database_restore', 0, 'RESTORE', new_values={'restored_from': 'Fictional Juniper verified backup'})
    restore_backup(backup.database_path, tmp_path / 'backups', audit=restore_audit)
    _assert_close(client_id)
    assert JournalEntry.count(client_id) == prior_count
    assert close_map.readiness(client_id, year.id)['ready']
    assert drafts.load(client_id)[0] == saved_revision
    assert all(BankReconciliation.get_by_id(i, client_id).status == 'Completed' for i in reconciliations)
    assert all('(AI)' not in (log.performed_by or '') for log in AuditLog.get_all(client_id))
    with dbconn.get_cursor() as cur:
        assert cur.execute('PRAGMA integrity_check').fetchone()[0] == 'ok'
        assert cur.execute('SELECT COUNT(*) FROM audit_log WHERE action="RESTORE"').fetchone()[0] == 1
        assert not cur.execute('SELECT id FROM journal_entries WHERE id IN (SELECT journal_entry_id FROM journal_entry_lines GROUP BY journal_entry_id HAVING SUM(debit) != SUM(credit))').fetchall()
    (tmp_path / 'verified-close.json').write_text(json.dumps({
        'synthetic': True, 'source_rows': 250, 'posted_imports': 244,
        'journal_entries': 248, 'trial_balance_cents': 1760000,
        'net_income_cents': 752500, 'expected_balances_cents': BALANCES,
        'reconciled_accounts': 2, 'backup_restored': True, 'integrity_check': 'ok',
    }, indent=2))
