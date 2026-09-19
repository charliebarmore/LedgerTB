"""Pure worksheet workbook rendering, shared by the session export cache."""

from datetime import datetime
from io import BytesIO
from utils.export import set_excel_literal


def build_worksheet_excel(client_name, period_start, period_end, rows, py_comparison):
    output = BytesIO()
    py_by_account = {row["account_number"]: row for row in py_comparison["accounts"]}
    current_account_numbers = {row.account_number for row in rows}
    import openpyxl
    from openpyxl.styles import Font, Alignment, Border, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Trial Balance Worksheet"

    # Header
    set_excel_literal(ws["A1"], f"Trial Balance Worksheet - {client_name}")
    ws["A1"].font = Font(bold=True, size=14)
    set_excel_literal(
        ws["A2"],
        f"Period: {period_start.strftime('%m/%d/%Y')} - {period_end.strftime('%m/%d/%Y')}",
    )
    set_excel_literal(
        ws["A3"],
        f"Generated: {datetime.now().strftime('%m/%d/%Y %H:%M')}",
    )

    # Column headers starting at row 5
    headers = [
        "Acct #",
        "Account Name",
        "Type",
        "Beg Bal Dr",
        "Beg Bal Cr",
        "Debits",
        "Credits",
        "Unadj TB Dr",
        "Unadj TB Cr",
        "AJE Dr",
        "AJE Cr",
        "Adj TB Dr",
        "Adj TB Cr",
        "PY Final Dr",
        "PY Final Cr",
    ]

    for col_idx, header in enumerate(headers, 1):
        cell = set_excel_literal(ws.cell(row=5, column=col_idx), header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    # Data rows starting at row 6. Keep accounts that existed only in
    # PY (for example a closed P&L account) visible with blank current
    # columns rather than silently dropping their comparison.
    data_start_row = 6
    export_rows = [(row, py_by_account.get(row.account_number, {})) for row in rows]
    export_rows.extend(
        (None, comparison_row)
        for comparison_row in py_comparison["accounts"]
        if comparison_row["account_number"] not in current_account_numbers
        and (comparison_row.get("prior_debit") or comparison_row.get("prior_credit"))
    )
    for row_idx, (row, comparison_row) in enumerate(export_rows, data_start_row):
        set_excel_literal(
            ws.cell(row=row_idx, column=1),
            row.account_number if row else comparison_row["account_number"],
        )
        set_excel_literal(
            ws.cell(row=row_idx, column=2),
            row.account_name if row else comparison_row["name"],
        )
        set_excel_literal(
            ws.cell(row=row_idx, column=3),
            row.account_type if row else comparison_row["type"],
        )
        ws.cell(
            row=row_idx,
            column=4,
            value=row.beginning_dr if row and row.beginning_dr > 0 else None,
        )
        ws.cell(
            row=row_idx,
            column=5,
            value=row.beginning_cr if row and row.beginning_cr > 0 else None,
        )
        ws.cell(
            row=row_idx,
            column=6,
            value=row.period_debits if row and row.period_debits > 0 else None,
        )
        ws.cell(
            row=row_idx,
            column=7,
            value=row.period_credits if row and row.period_credits > 0 else None,
        )
        # Unadjusted TB uses formulas
        ws.cell(
            row=row_idx,
            column=8,
            value=f"=MAX(D{row_idx}-E{row_idx}+F{row_idx}-G{row_idx},0)",
        )
        ws.cell(
            row=row_idx,
            column=9,
            value=f"=MAX(E{row_idx}-D{row_idx}+G{row_idx}-F{row_idx},0)",
        )
        ws.cell(
            row=row_idx,
            column=10,
            value=row.aje_debits if row and row.aje_debits > 0 else None,
        )
        ws.cell(
            row=row_idx,
            column=11,
            value=row.aje_credits if row and row.aje_credits > 0 else None,
        )
        # Adjusted TB uses formulas
        ws.cell(
            row=row_idx,
            column=12,
            value=f"=MAX(H{row_idx}-I{row_idx}+J{row_idx}-K{row_idx},0)",
        )
        ws.cell(
            row=row_idx,
            column=13,
            value=f"=MAX(I{row_idx}-H{row_idx}+K{row_idx}-J{row_idx},0)",
        )
        ws.cell(row=row_idx, column=14, value=comparison_row.get("prior_debit") or None)
        ws.cell(
            row=row_idx, column=15, value=comparison_row.get("prior_credit") or None
        )

    # Totals row with formulas
    totals_row = data_start_row + len(export_rows)
    set_excel_literal(ws.cell(row=totals_row, column=1), "TOTALS")
    ws.cell(row=totals_row, column=1).font = Font(bold=True)

    for col_idx in range(4, 16):
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws.cell(
            row=totals_row,
            column=col_idx,
            value=f"=SUM({col_letter}{data_start_row}:{col_letter}{totals_row-1})",
        )
        ws.cell(row=totals_row, column=col_idx).font = Font(bold=True)

    # Format number columns
    for row in ws.iter_rows(
        min_row=data_start_row, max_row=totals_row, min_col=4, max_col=15
    ):
        for cell in row:
            cell.number_format = "#,##0.00"

    # Adjust column widths
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 10
    for col in ["D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N", "O"]:
        ws.column_dimensions[col].width = 12

    wb.save(output)
    return output.getvalue()
